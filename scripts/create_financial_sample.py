"""Create synthetic financial test data without callbacks or real account details."""
from pathlib import Path
from docx import Document

LAB_FINANCIAL_RECORD_ID = 'LAB-FIN-20270910'

def create_sample(output: Path) -> Path:
    output.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading('Financial Records', 0)
    document.add_paragraph('CONFIDENTIAL')
    table = document.add_table(rows=1, cols=3)
    table.style = 'Light Shading Accent 1'
    for cell, value in zip(table.rows[0].cells, ['Period', 'Revenue', 'Operating costs']):
        cell.text = value
    for values in [('2027 Q1', '$125,000', '$82,000'), ('2027 Q2', '$132,000', '$86,000')]:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    document.add_paragraph('Lab financial record')
    document.add_paragraph(f'Record ID: {LAB_FINANCIAL_RECORD_ID}')
    document.add_paragraph('All financial data used in this lab is synthetic test data.')
    target = output / 'Financial_Records_CONFIDENTIAL.docx'
    document.save(target)
    return target

if __name__ == '__main__':
    print(create_sample(Path('generated')))
